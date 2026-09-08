"""Generate the D4.2 manual evidence form. Use the bundled document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "quality/manual-tests/ARTestCLI_Stage_D4_2_Manual_Test_Report.docx"
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles["Normal"].font.size = Pt(10)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
doc.styles["Title"].font.size = Pt(24)
doc.styles["Heading 1"].font.size = Pt(17)
doc.styles["Heading 2"].font.size = Pt(12)
section.header.paragraphs[0].text = "ARTestCLI   |   D4.2 Python host and SDK   |   Manual acceptance"
section.header.paragraphs[0].runs[0].font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.text = "Evidence status: NOT RUN until completed by the tester                    Page "
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
footer.runs[0].font.size = Pt(8)

def p(text, style=None): return doc.add_paragraph(text, style)
def code(text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    for line in text.splitlines():
        if paragraph.runs: paragraph.add_run("\n")
        run = paragraph.add_run(line)
        run.font.name, run.font.size = "Consolas", Pt(9)
    return paragraph
def table(rows, widths=None):
    result = doc.add_table(rows=0, cols=len(rows[0]))
    result.style = "Table Grid"
    result.autofit = False
    for values in rows:
        cells = result.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            if widths: cells[i].width = Inches(widths[i])
        trpr = result.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
    return result

doc.add_heading("ARTest D4 Python Manual Test Report", 0)
p("D4.2 Python host and SDK acceptance")
p("Use this report to verify the prepared Python environment, generated metadata, mixed C++ and "
  "Python execution, measurement verdicts, error handling and cleanup. All devices in these cases "
  "are simulated. Do not connect physical equipment.")
p("Automated execution does not complete this report. Record the actual behavior, attach screenshots "
  "or logs, and assign PASS or FAIL only after comparing every expected result.")
table([
    ("Field", "Tester entry"),
    ("Tester and execution date", "________________________________________"),
    ("Git commit or working tree identifier", "________________________________________"),
    ("Windows version and architecture", "________________________________________"),
    ("Visual Studio version and build configuration", "________________________________________"),
    ("Python path and version", "________________________________________"),
    ("Generated manual case directory", "________________________________________"),
    ("Overall verdict and open defect references", "NOT RUN"),
], [2.8, 4.2])
doc.add_heading("Acceptance rules", 1)
p("A negative scenario passes when ARTest reports the expected failure accurately. For example, "
  "MT03 requires a failed measurement and process exit code 5. Do not mark MT03 failed merely "
  "because the test sequence correctly ends FAILED.")
p("Never treat a killed worker as proof that physical hardware was cleaned up. The fault-injection "
  "case requires an indeterminate result, one attempt and an explicit cleanup warning.")
p("Before signing acceptance, complete MT01 through MT10 and reference the automated XML/HTML reports. "
  "Record any deviation as a defect rather than editing the expected results.")

doc.add_page_break()
doc.add_heading("Preparation and shared variables", 1)
p("Use a normal PowerShell terminal, not the Visual Studio Package Manager Console. "
  "Keep the same terminal open for all cases. If using another repository/interpreter path, "
  "record that change on the first page and update the commands consistently.")
code(r"""Set-Location 'D:\GitHub\main\ARTestCLI'
$python = 'C:\Users\anton\AppData\Local\Programs\Python\Python313\python.exe'
& $python -I -c "import sys; print(sys.version)"
.\scripts\build.ps1 -Configuration Release -Platform x64
.\scripts\prepare-python-example.ps1 -Python $python -IncludeFaultTests
$cli = '.\artifacts\bin\x64\Release\ARTestCLI.exe'
$mapping = '.\artifacts\python\environments\python-environments.json'
$cases = .\scripts\manual-tests\prepare-stage-d4-2.ps1 -Configuration Release
$catalog = Join-Path $cases 'catalog'
$cases""")
p("Expected prerequisites: standard CPython 3.13 x64; successful Release build; prepared environment "
  "mapping; a new case directory. Preparation uses pinned protobuf 6.33.4 and pywin32 311 in isolated "
  "venvs. No global pip installation is required. If changed inputs cause a preparation mismatch, "
  "stop and prepare a new output root as described in the Python guide; do not edit receipts.")
doc.add_heading("Evidence capture", 2)
p("For each execution case, run the command shown and then $LASTEXITCODE immediately. Capture the "
  "terminal with the final JSON and exit code. The JSON is printed to the terminal; it is not "
  "automatically saved to a report file. Use screenshots or explicitly redirect/copy output.")
p("All generated plans and copied catalogs are under $cases. The preparation script creates a "
  "unique directory and does not overwrite previous manual evidence.")
p("Reference: docs/sdk/python-extension-authoring.md. Architecture and deferred planning findings: "
  "docs/architecture/stage-d4-2-python-runtime.md.")

cases = [
("MT01", "Generated metadata and offline compilation", None,
 r"""Get-ChildItem '.\artifacts\python\extensions\ARTestPySimulated'
Get-Content '.\artifacts\python\extensions\ARTestPySimulated\artest-extension.json'
& $cli compile "$cases\PythonPassed.json" --extensions $catalog
$LASTEXITCODE""",
 ["The package contains code, schemas, requirements.lock and generated artest-extension.json.",
  "The manifest declares schemaVersion 3, runtime kind python, outOfProcess, runtimeVersion 3.13 and a file inventory.",
  "Compilation succeeds with exit code 0 and states that no instruments were initialized.",
  "No PYTHON_WORKER_READY or PY_DRIVER_INITIALIZE message appears during compilation."]),
("MT02", "Python command and Python driver", "PythonPassed", None,
 ["Exit code 0; overall status and step status are passed.",
  "Events include PYTHON_WORKER_READY, PY_DRIVER_INITIALIZE, PY_MEASUREMENT_COMPLETED and PY_DRIVER_SHUTDOWN.",
  "Final schema is artest.schema.run-result.v2. The first step outcome contains value 5.0, unit V and minimum 4.8.",
  "Summary reports one planned, executed and passed step."]),
("MT03", "Failed measurement remains a failed verdict", "PythonFailed", None,
 ["Exit code 5; overall and step status are failed, not passed or error.",
  "Summary has failedSteps 1 and errorSteps 0.",
  "Step and attempt outcome data retain value 4.2, minimum 4.8 and unit V.",
  "PY_DRIVER_SHUTDOWN appears before the terminal state."]),
("MT04", "Native command calls Python driver", "NativeToPython", None,
 ["Exit code 0; overall and step status are passed.",
  "The command is com.artest.command.sample.power-cycle and it uses the configured Python driver.",
  "PY_DRIVER_INITIALIZE and PY_DRIVER_SHUTDOWN appear. The native sample command reports service completion.",
  "No direct command-to-driver DLL dependency or manual JSON metadata edit is needed."]),
("MT05", "Python command calls native driver", "PythonToNative", None,
 ["Exit code 0; overall and step status are passed.",
  "The native simulated power driver reports initialization and shutdown.",
  "PY_MEASUREMENT_COMPLETED appears; the outcome retains measured value 5.0 and unit V.",
  "The Python command implementation is unchanged from MT02."]),
("MT06", "Python exception becomes an execution error", "PythonError", None,
 ["Exit code 5; overall status is error and summary errorSteps is 1.",
  "The diagnostic includes Simulated command error.",
  "The exception does not terminate ARTestCLI abnormally.",
  "PY_DRIVER_SHUTDOWN appears and the next independent run of MT02 can pass."]),
("MT07", "Deadline and cooperative cleanup", "PythonTimeout", None,
 ["Exit code 5; overall and step status are timedOut.",
  "The 100 ms step deadline interrupts a requested 2000 ms hold.",
  "PY_DRIVER_SHUTDOWN appears. The step outcome indeterminate flag is false.",
  "Do not use total process duration as the step latency; startup/environment validation are separate."]),
("MT08", "User cancellation and cleanup", "PythonCancel", None,
 ["Wait until the terminal prints Executing step 1, then press Ctrl+C once within 5 seconds.",
  "ARTest handles cancellation and returns exit code 5; overall status is cancelled.",
  "PY_DRIVER_SHUTDOWN appears and the step outcome indeterminate flag is false.",
  "The terminal remains usable. Repeating MT02 in the same terminal succeeds."]),
("MT09", "Cleanup failure overrides successful measurement", "PythonCleanupError", None,
 ["The step measurement passes, but the overall run status is error and exit code is 5.",
  "The diagnostic includes Simulated cleanup failure.",
  "The run must not report overall PASSED merely because the measurement passed.",
  "PY_DRIVER_SHUTDOWN confirms the cleanup attempt, not its successful completion."]),
("MT10", "Controlled process faults and no replay", None,
 r""".\scripts\test-python-runtime.ps1 -Configuration Release
Get-ChildItem '.\artifacts\test-results\x64\Release\ARTestPython.Integration.*'""",
 ["The script succeeds and reports all Python integration cases passed.",
  "Open ARTestPython.Integration.html and verify CrashAfterEffectIsNeverRetriedOrContinued, "
  "BlockingIoThatIgnoresTimeoutIsTerminatedWithoutReplay and BlockingIoFinishesBeforeDriverCleanupBegins are PASSED.",
  "The fault tests verify one persisted simulated effect, no retry/continuation, explicit indeterminate state and unconfirmed cleanup after worker termination.",
  "Attach the HTML/XML report references and a screenshot. No real device or Task Manager kill is required."])
]
for identity, title, plan, command, expected in cases:
    doc.add_page_break()
    doc.add_heading(identity + " " + title, 1)
    p("Prerequisite: complete the preparation page and retain its PowerShell variables.")
    doc.add_heading("Procedure", 2)
    if identity == "MT08":
        p("Run the command below. After Executing step 1 appears, press Ctrl+C once within 5 seconds. "
          "Wait for ARTest to finish cleanup, then run $LASTEXITCODE.")
    if command is None:
        command = '& $cli extension-run "$cases\\' + plan + '.json" $catalog --python-environments $mapping\n$LASTEXITCODE'
    code(command)
    doc.add_heading("Expected results", 2)
    for index, text in enumerate(expected, 1): p(str(index) + ". " + text)
    doc.add_heading("Actual results and evidence", 2)
    table([
        ("Verdict", "NOT RUN   /   PASS   /   FAIL   /   BLOCKED"),
        ("Actual result and exit code", "\n\n"),
        ("Evidence file or screenshot", "\n\n\n"),
        ("Defect reference and notes", "\n"),
        ("Tester and date", ""),
    ], [2.05, 4.95])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
for tree in (doc.styles.element, doc.element):
    for border in tree.xpath(".//w:pBdr"):
        border.getparent().remove(border)
doc.save(OUTPUT)
print(OUTPUT)
