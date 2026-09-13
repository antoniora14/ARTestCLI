"""Create the unexecuted C-02 evidence form with the bundled document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "quality/manual-tests/checkpoint-c02/ARTest_C02_Manual_Test_Report.docx"
if OUTPUT.exists():
    raise FileExistsError("Preserve tester evidence; select a new report revision instead.")
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.75)
for name, size in (("Normal", 11), ("Title", 23), ("Heading 1", 17), ("Heading 2", 12)):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
footer = section.footer.paragraphs[0]
footer.text = "ARTest C02  |  Manual evidence  |  Page "
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
footer.runs[0].font.size = Pt(9)

def paragraph(text):
    return doc.add_paragraph(text)

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    for index, line in enumerate(text.splitlines()):
        if index:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name, run.font.size = "Consolas", Pt(9)
    return p

def evidence():
    doc.add_heading("Execution record", 2)
    paragraph("Verdict: NOT RUN    Tester: __________________    Date: ______________")
    paragraph("Actual result and defect reference: __________________________________")
    paragraph("Evidence path or screenshot caption: _________________________________")
    paragraph("Attach screenshots and logs below or on additional pages.")
    paragraph("\n\n")

doc.add_heading("ARTest C02 Manual Acceptance Report", 0)
paragraph("TCP driver and native and Python command acceptance")
paragraph("Verify the independent SDK example, two instrument instances, generated metadata "
          "and correct reporting of failed measurements and uncertain external effects. "
          "Use only the loopback simulator. No physical instrument is required.")
paragraph("Automated results do not complete this form. Record the observed behavior and "
          "attach evidence before replacing NOT RUN with PASS or FAIL. A negative scenario "
          "passes when the expected failure is reported accurately.")
for text in ("Tester and execution date: __________________________________________",
             "Commit and working tree identifier: __________________________________",
             "Windows and Visual Studio versions: __________________________________",
             "Configuration and native SDK ZIP SHA256: ______________________________",
             "Python executable and version: _______________________________________",
             "Overall verdict and open defects: NOT RUN"):
    paragraph(text)
doc.add_heading("Shared preparation", 1)
paragraph("Use PowerShell 7 and keep this terminal open. Release is used throughout. "
          "Run the matching full build first if binaries are absent. Do not change historical "
          "SDK baselines or pending reports. These variables use the current repository path.")
code(r"""$repo = 'D:\GitHub\main\ARTestCLI'
Set-Location $repo
$cli = "$repo\artifacts\bin\x64\Release\ARTestCLI.exe"
$sdk = "$repo\artifacts\sdk-packages\x64\Release"
$zip = "$sdk\ARTestSDK-0.4.0-windows-x64.zip"
$example = "$repo\examples\ARTestTcpHello"
$python = "$env:LOCALAPPDATA\Programs\Python\Python313\python.exe"
Test-Path $cli
Test-Path $zip
& $python --version""")
paragraph("Expected: both paths exist and Python reports 3.13.x x64. Python is needed only "
          "for MT03 and the Python part of MT04. Record actual tool versions above.")

cases = [
("MT01 Independent copied kit", [
    "Run the external-consumer gate below. It copies source into a unique TEMP directory, "
    "extracts the SDK ZIP, checks its inventory, builds the DLL and simulator, compiles "
    "offline, and executes two native measurements.",
    "Record the externalRoot and receipt.json paths printed at completion. Open the copied "
    "folder in File Explorer and inspect the generated package under out/extensions/x64/Release."],
 r""".\scripts\test-tcp-hello-external.ps1 -Configuration Release -SDKArchive $zip""",
 "The gate finishes successfully. The copied project builds without repository headers or "
 "Engine/Core project links. The generated package contains the DLL, artest-extension.json "
 "and schemas. The final result has two passed steps, values 12 V and 5 V, and the server "
 "journal ends with stopped. Attach build output, the receipt and the package screenshot."),
("MT02 Native execution and offline compilation", [
    "Do not start the simulator yourself. First compile the source plan offline, then run "
    "the example. Run.ps1 obtains an OS-assigned port and stops its own simulator.",
    "Open the printed run directory. Inspect plan.json, cli.txt and server.jsonl. Repeat "
    "Run.ps1 once to verify a later session starts cleanly."],
 r"""& $cli compile "$example\ExamplePlan.json" --extensions "$example\out\extensions\x64\Release"
& "$example\Run.ps1" -CLI $cli -Configuration Release""",
 "Offline compilation succeeds without initializing instruments. Each run reports two passed "
 "steps with instrumentId V1/value 12 and V2/value 5, schema artest.schema.example.tcp-measurement.v1. "
 "Each journal contains exactly two applied events and two close requests. Verify the process "
 "identified by ready.json has exited; do not terminate unrelated processes."),
("MT03 Python command and native driver", [
    "Prepare the explicit Python environment using the existing C-01 SDK wheel. This is "
    "separate from the C-01 accepted environment and does not modify it.",
    "Execute the same example with the Python package and environment receipt. Inspect "
    "the final JSON and the simulator-side journal."],
 r"""$wheel = "$repo\artifacts\python-c01-final\sdk"
$wheel = "$wheel\artest_python-0.2.0-py3-none-any.whl"
.\scripts\prepare-tcp-hello-python.ps1 -Python $python -SDKWheel $wheel
$py = "$repo\artifacts\python-c02"
& "$example\Run.ps1" -CLI $cli -PythonPackage "$py\extensions\ARTestPyTcpHello" -PythonEnvironment "$py\environment\artest-environment.json" """,
 "Preparation passes dependency and integrity checks. Both Python command steps pass, "
 "returning the same measurement schema and values as MT02. The native simulator records "
 "exactly two applied actions. Cleanup is logged and the owned simulator exits. "
 "A successful native run alone is not evidence that this Python path passed."),
("MT04 Failure verdicts and uncertainty evidence", [
    "Run the focused gate below in a new evidence directory. The fault modes belong only "
    "to the test server process, never to the normal simulator.",
    "Open tests.html and the session JSON files under servers. Find a failed-limit session "
    "and a session whose first outcome is indeterminate. Match that session's plan port to "
    "a server ready.json, then inspect its mode.json and journal.jsonl."],
 r"""$evidence = "$repo\artifacts\manual-c02-" + [guid]::NewGuid().ToString('N')
.\scripts\test-tcp-hello.ps1 -Configuration Release -IncludePython -OutputDirectory $evidence""",
 "Exactly 20 cases execute and pass, with no skipped cases counted as passes. The failed-limit "
 "session is failed rather than passed. Lost ACK has one applied journal event, one attempt, "
 "one skipped next step and indeterminate=true on step and attempt despite maxAttempts=3 "
 "and onFailure=continue. Cleanup is attempted. Cancellation/timeout stays on step and attempt; "
 "global error also reflects the unconfirmed close. The pre-send negative control has zero "
 "applied events and six attempts across two steps."),
]
for title, steps, command, expected in cases:
    doc.add_page_break()
    doc.add_heading(title, 1)
    doc.add_heading("Procedure", 2)
    for index, text in enumerate(steps, 1):
        paragraph(f"{index}. {text}")
    code(command.strip())
    doc.add_heading("Expected results", 2)
    paragraph(expected)
    evidence()
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
